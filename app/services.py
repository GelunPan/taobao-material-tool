"""业务服务层：图片处理与 Excel 导出等可复用能力。

与界面解耦：服务只接收数据与路径，不直接操作窗口控件。
"""
import os
from pathlib import Path

from PyQt6.QtGui import QImage

from .config import EXPORT_COLUMN_WIDTHS, RECORD_FIELDS, TABLE_HEADERS, IMAGE_EXT

# 复制图片文件粘贴时识别的扩展名
_IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".webp", ".tif", ".tiff", ".ico"}


class ImageService:
    """图片相关服务：剪贴板图片保存、文件名生成"""

    @staticmethod
    def save_clipboard_image(clipboard, images_dir: Path, counter: int):
        """保存剪贴板中的图片，支持两种来源：

        1. 剪贴板二进制图片（截图 / 复制图片内容）→ 保存到 images_dir；
        2. 复制的图片文件（文件管理器 Ctrl+C）→ 直接引用原文件路径。

        返回 (文件路径, 新计数器)；剪贴板无有效图片时返回 (None, counter)。
        """
        mime = clipboard.mimeData()
        if mime is None:
            return None, counter

        # 来源一：剪贴板二进制图片（截图等）
        image = ImageService._image_from_mime(clipboard, mime)
        if not image.isNull():
            counter += 1
            images_dir.mkdir(parents=True, exist_ok=True)
            filepath = images_dir / f"image_{counter:04d}.{IMAGE_EXT.lower()}"
            if image.save(str(filepath), IMAGE_EXT):
                return str(filepath), counter
            return None, counter - 1

        # 来源二：复制的图片文件（text/uri-list，文件管理器复制文件场景）
        path = ImageService._first_image_file(mime)
        if path:
            return path, counter

        return None, counter

    @staticmethod
    def _image_from_mime(clipboard, mime):
        """从剪贴板提取图片 QImage：优先 clipboard.image()，失败再从原始字节兜底解码。

        hasImage() 只认 Qt 内部格式，部分程序复制图片时仅写入 image/png 等字节，
        此时 hasImage() 为 False，因此这里不依赖 hasImage() 而直接尝试字节解码。
        """
        if mime.hasImage():
            image = clipboard.image()
            if not image.isNull():
                return image
        return ImageService._decode_from_mime(mime)

    @staticmethod
    def _first_image_file(mime) -> str | None:
        """从剪贴板的文件 URL 里取第一个本地图片文件路径，无则返回 None"""
        if not mime.hasUrls():
            return None
        for url in mime.urls():
            local = url.toLocalFile()
            if local and os.path.exists(local) and Path(local).suffix.lower() in _IMAGE_SUFFIXES:
                return local
        return None

    @staticmethod
    def _decode_from_mime(mime):
        """从剪贴板 MIME 数据中按常见图片格式兜底解码出 QImage，失败返回 null"""
        image = QImage()
        for fmt in ("image/png", "image/bmp", "image/jpeg", "image/webp",
                    "image/gif", "image/tiff"):
            if mime.hasFormat(fmt) and image.loadFromData(mime.data(fmt)):
                return image
        return QImage()


class ExcelExporter:
    """将素材记录导出为 Excel 文件"""

    @staticmethod
    def export(records: list, file_path: str, sheet_name: str = "素材") -> None:
        from openpyxl import Workbook
        from openpyxl.styles import Font

        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name

        # 表头
        for col, header in enumerate(TABLE_HEADERS, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)

        # 数据行（字段顺序与列顺序一一对应；多图列表换行拼接）
        for row, record in enumerate(records, 2):
            for col, field in enumerate(RECORD_FIELDS, 1):
                value = record.get(field, "")
                if isinstance(value, list):
                    value = "\n".join(str(v) for v in value)
                ws.cell(row=row, column=col, value=value)

        # 列宽
        for col, width in enumerate(EXPORT_COLUMN_WIDTHS, 1):
            ws.column_dimensions[chr(64 + col)].width = width

        wb.save(file_path)
